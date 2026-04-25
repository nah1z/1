"""
產生情侶共同基金記帳 Excel 範本（.xlsx）

五個工作表：
- 基金清單：定義有哪些基金（提供下拉選單來源）
- 預算：每個基金每月可累加多筆預算
- 交易：所有收支紀錄
- 月度報表：選基金與月份後自動算出當月預算 / 收入 / 支出 / 結餘
- 快速記帳：填輸入區 → 按巨集按鈕 → 自動 append 到「交易」/「預算」
            （需另存為 .xlsm 並匯入 scripts/macros/CoupleBudget.bas）
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = "/home/user/1/couple-budget-template.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
SECTION_FONT = Font(bold=True, size=11, color="1F4E78")
INCOME_FILL = PatternFill("solid", fgColor="E2EFDA")
EXPENSE_FILL = PatternFill("solid", fgColor="FCE4D6")
BUDGET_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

MONEY_FMT = '_-"$"* #,##0.00_-;-"$"* #,##0.00_-;_-"$"* "-"??_-;_-@_-'
DATE_FMT = "yyyy-mm-dd"
MONTH_FMT = "yyyy-mm"

# 預先填入的範例資料量（可放空，供使用者參考格式）
N_FUNDS = 5            # 基金清單可填 5 列
N_BUDGETS = 200        # 預算最多 200 列
N_TRANSACTIONS = 1000  # 交易最多 1000 列


def style_header(ws, row, col_count):
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def style_data_range(ws, top, bottom, col_count):
    for r in range(top, bottom + 1):
        for c in range(1, col_count + 1):
            ws.cell(row=r, column=c).border = BORDER


def make_funds(ws):
    ws.title = "基金清單"
    ws["A1"] = "基金清單"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:C1")

    ws["A2"] = "在這裡定義所有共同基金。基金名稱會被「預算」「交易」「月度報表」工作表引用。"
    ws["A2"].font = Font(italic=True, color="595959")
    ws.merge_cells("A2:C2")

    headers = ["基金名稱", "建立日期", "備註"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))

    samples = [
        ("日常", "2026-04-01", "兩人每月共同生活開銷"),
        ("旅遊", "2026-04-01", "存旅費用"),
    ]
    for r, row in enumerate(samples, start=5):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)

    style_data_range(ws, 5, 4 + N_FUNDS, len(headers))
    for r in range(5, 5 + N_FUNDS):
        ws.cell(row=r, column=2).number_format = DATE_FMT

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 40


def make_budgets(ws):
    ws.title = "預算"
    ws["A1"] = "預算（同基金同月份可多筆，總額為加總）"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:E1")

    ws["A2"] = "每月初放入預算記在這裡。月份格式為 YYYY-MM，例如 2026-04。"
    ws["A2"].font = Font(italic=True, color="595959")
    ws.merge_cells("A2:E2")

    headers = ["基金", "月份 (YYYY-MM)", "金額", "記錄人", "備註"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))

    samples = [
        ("日常", "2026-04", 15000, "A", "月初預算"),
        ("日常", "2026-04", 5000, "B", "中旬補充"),
        ("旅遊", "2026-04", 8000, "A", "京都計畫"),
    ]
    for r, row in enumerate(samples, start=5):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)

    bottom = 4 + N_BUDGETS
    style_data_range(ws, 5, bottom, len(headers))
    for r in range(5, bottom + 1):
        ws.cell(row=r, column=3).number_format = MONEY_FMT
        ws.cell(row=r, column=3).fill = BUDGET_FILL

    widths = [16, 18, 14, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def make_transactions(ws):
    ws.title = "交易"
    ws["A1"] = "交易紀錄（所有收支）"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:G1")

    ws["A2"] = "每一筆收入或支出記在這裡。類型欄會自動以下拉選單限定為「收入」或「支出」。"
    ws["A2"].font = Font(italic=True, color="595959")
    ws.merge_cells("A2:G2")

    headers = ["日期", "基金", "類型", "項目", "金額", "記錄人", "備註"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))

    samples = [
        ("2026-04-02", "日常", "支出", "超市", 1280, "A", ""),
        ("2026-04-03", "日常", "支出", "晚餐", 850, "B", "義式"),
        ("2026-04-05", "日常", "收入", "退款", 200, "A", "退超市"),
        ("2026-04-10", "旅遊", "支出", "訂機票", 6500, "A", "華航"),
    ]
    for r, row in enumerate(samples, start=5):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)

    bottom = 4 + N_TRANSACTIONS
    style_data_range(ws, 5, bottom, len(headers))
    for r in range(5, bottom + 1):
        ws.cell(row=r, column=1).number_format = DATE_FMT
        ws.cell(row=r, column=5).number_format = MONEY_FMT

    widths = [12, 14, 10, 24, 14, 12, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def make_report(ws):
    ws.title = "月度報表"
    ws["A1"] = "月度報表"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    ws["A2"] = "選擇基金與月份，下方統計會自動更新。"
    ws["A2"].font = Font(italic=True, color="595959")
    ws.merge_cells("A2:D2")

    ws["A4"] = "選擇基金"
    ws["A4"].font = SECTION_FONT
    ws["B4"] = "日常"
    ws["B4"].alignment = LEFT
    ws["B4"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["B4"].border = BORDER

    ws["A5"] = "選擇月份"
    ws["A5"].font = SECTION_FONT
    ws["B5"] = "2026-04"
    ws["B5"].alignment = LEFT
    ws["B5"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["B5"].border = BORDER

    ws["A7"] = "本月統計"
    ws["A7"].font = SECTION_FONT

    ws["A8"] = "本月預算"
    ws["B8"] = (
        '=SUMIFS(預算!C:C,預算!A:A,$B$4,預算!B:B,$B$5)'
    )
    ws["B8"].fill = BUDGET_FILL

    ws["A9"] = "本月收入"
    ws["B9"] = (
        '=SUMIFS(交易!E:E,交易!B:B,$B$4,交易!C:C,"收入",'
        '交易!A:A,">="&DATE(VALUE(LEFT($B$5,4)),VALUE(RIGHT($B$5,2)),1),'
        '交易!A:A,"<"&EDATE(DATE(VALUE(LEFT($B$5,4)),VALUE(RIGHT($B$5,2)),1),1))'
    )
    ws["B9"].fill = INCOME_FILL

    ws["A10"] = "本月支出"
    ws["B10"] = (
        '=SUMIFS(交易!E:E,交易!B:B,$B$4,交易!C:C,"支出",'
        '交易!A:A,">="&DATE(VALUE(LEFT($B$5,4)),VALUE(RIGHT($B$5,2)),1),'
        '交易!A:A,"<"&EDATE(DATE(VALUE(LEFT($B$5,4)),VALUE(RIGHT($B$5,2)),1),1))'
    )
    ws["B10"].fill = EXPENSE_FILL

    ws["A11"] = "結餘"
    ws["B11"] = "=B8+B9-B10"
    ws["B11"].font = Font(bold=True, size=12)

    for r in range(8, 12):
        ws.cell(row=r, column=2).number_format = MONEY_FMT
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=2).border = BORDER

    ws["A13"] = "本月交易明細（依日期）"
    ws["A13"].font = SECTION_FONT

    headers = ["日期", "類型", "項目", "金額", "記錄人", "備註"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=14, column=i, value=h)
    style_header(ws, 14, len(headers))

    ws["A15"] = (
        '=IFERROR(SORT(FILTER(CHOOSE({1,2,3,4,5,6},'
        '交易!A5:A1004,交易!C5:C1004,交易!D5:D1004,'
        '交易!E5:E1004,交易!F5:F1004,交易!G5:G1004),'
        '(交易!B5:B1004=$B$4)*'
        '(TEXT(交易!A5:A1004,"yyyy-mm")=$B$5)),1,1),"")'
    )

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 30

    for r in range(15, 60):
        ws.cell(row=r, column=1).number_format = DATE_FMT
        ws.cell(row=r, column=4).number_format = MONEY_FMT


def make_quickadd(ws):
    ws.title = "快速記帳"
    ws["A1"] = "快速記帳（巨集輔助）"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")

    ws["A2"] = (
        "填好輸入區後按對應按鈕（需先把檔案另存為 .xlsm 並匯入 CoupleBudget.bas，"
        "詳見下方步驟）。"
    )
    ws["A2"].font = Font(italic=True, color="595959")
    ws.merge_cells("A2:F2")

    INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")

    # 左側：交易輸入
    ws["A3"] = "新增交易"
    ws["A3"].font = SECTION_FONT
    tx_fields = [
        ("日期", "2026-04-15", DATE_FMT),
        ("基金", "日常", None),
        ("類型", "支出", None),
        ("項目", "", None),
        ("金額", "", MONEY_FMT),
        ("記錄人", "A", None),
        ("備註", "", None),
    ]
    for i, (label, default, fmt) in enumerate(tx_fields):
        r = 4 + i
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=1).alignment = LEFT
        ws.cell(row=r, column=1).border = BORDER
        v = ws.cell(row=r, column=2, value=default)
        v.fill = INPUT_FILL
        v.border = BORDER
        v.alignment = LEFT
        if fmt:
            v.number_format = fmt

    ws["A12"] = "→ 在此處附近插入按鈕並指派巨集 AddTransaction"
    ws["A12"].font = Font(italic=True, color="C00000")
    ws.merge_cells("A12:B12")

    # 右側：預算輸入
    ws["D3"] = "新增預算"
    ws["D3"].font = SECTION_FONT
    bg_fields = [
        ("基金", "日常", None),
        ("月份 (YYYY-MM)", "2026-04", None),
        ("金額", "", MONEY_FMT),
        ("記錄人", "A", None),
        ("備註", "", None),
    ]
    for i, (label, default, fmt) in enumerate(bg_fields):
        r = 4 + i
        ws.cell(row=r, column=4, value=label).font = Font(bold=True)
        ws.cell(row=r, column=4).alignment = LEFT
        ws.cell(row=r, column=4).border = BORDER
        v = ws.cell(row=r, column=5, value=default)
        v.fill = INPUT_FILL
        v.border = BORDER
        v.alignment = LEFT
        if fmt:
            v.number_format = fmt

    ws["D10"] = "→ 在此處附近插入按鈕並指派巨集 AddBudget"
    ws["D10"].font = Font(italic=True, color="C00000")
    ws.merge_cells("D10:E10")

    # 巨集設定步驟
    ws["A14"] = "巨集設定步驟（一次性，約 2 分鐘）"
    ws["A14"].font = SECTION_FONT
    steps = [
        "1. 把這個 .xlsx 「另存新檔」 → 檔案類型選「Excel 啟用巨集的活頁簿 (.xlsm)」。",
        "2. 按 Alt + F11 開啟 VBA 編輯器。",
        "3. 上方選單 File → Import File... → 選擇 scripts/macros/CoupleBudget.bas。",
        "4. 關閉 VBA 編輯器，回到 Excel。",
        "5. 確認看到「開發人員」索引標籤；沒看到時：檔案 → 選項 → 自訂功能區 → 勾「開發人員」。",
        "6. 開發人員 → 插入 → 表單控制項「按鈕」→ 在交易輸入區下方拖出一個按鈕，跳出視窗時選 AddTransaction。",
        "7. 同樣步驟在預算輸入區下方插入按鈕，指派 AddBudget。可右鍵按鈕 → 編輯文字 改成「新增交易」/「新增預算」。",
        "8. 儲存檔案。之後填好黃色輸入區、按按鈕即可。",
    ]
    for i, s in enumerate(steps):
        r = 15 + i
        ws.cell(row=r, column=1, value=s)
        ws.merge_cells(start_row=r, end_row=r, start_column=1, end_column=6)

    widths = [22, 22, 4, 22, 22, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_validations(wb):
    fund_range = "=基金清單!$A$5:$A$" + str(4 + N_FUNDS)

    dv_fund = DataValidation(type="list", formula1=fund_range, allow_blank=True)
    dv_fund.error = "請從「基金清單」中選擇一個基金名稱"
    dv_fund.errorTitle = "無效的基金"
    wb["預算"].add_data_validation(dv_fund)
    dv_fund.add(f"A5:A{4 + N_BUDGETS}")

    dv_fund2 = DataValidation(type="list", formula1=fund_range, allow_blank=True)
    wb["交易"].add_data_validation(dv_fund2)
    dv_fund2.add(f"B5:B{4 + N_TRANSACTIONS}")

    dv_fund3 = DataValidation(type="list", formula1=fund_range, allow_blank=False)
    wb["月度報表"].add_data_validation(dv_fund3)
    dv_fund3.add("B4")

    dv_type = DataValidation(type="list", formula1='"收入,支出"', allow_blank=False)
    dv_type.error = "類型只能是「收入」或「支出」"
    dv_type.errorTitle = "無效的類型"
    wb["交易"].add_data_validation(dv_type)
    dv_type.add(f"C5:C{4 + N_TRANSACTIONS}")

    # 快速記帳輸入區的下拉
    qa = wb["快速記帳"]
    dv_qa_fund_tx = DataValidation(type="list", formula1=fund_range, allow_blank=False)
    qa.add_data_validation(dv_qa_fund_tx)
    dv_qa_fund_tx.add("B5")

    dv_qa_type = DataValidation(type="list", formula1='"收入,支出"', allow_blank=False)
    qa.add_data_validation(dv_qa_type)
    dv_qa_type.add("B6")

    dv_qa_fund_bg = DataValidation(type="list", formula1=fund_range, allow_blank=False)
    qa.add_data_validation(dv_qa_fund_bg)
    dv_qa_fund_bg.add("E4")


def main():
    wb = Workbook()
    make_funds(wb.active)
    make_budgets(wb.create_sheet())
    make_transactions(wb.create_sheet())
    make_report(wb.create_sheet())
    make_quickadd(wb.create_sheet())
    add_validations(wb)

    wb["快速記帳"].sheet_view.tabSelected = True
    wb.active = wb.index(wb["快速記帳"])

    wb.save(OUTPUT)
    print(f"saved: {OUTPUT}")


if __name__ == "__main__":
    main()
